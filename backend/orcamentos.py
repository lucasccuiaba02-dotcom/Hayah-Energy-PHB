"""Set/2026 · Orçamentos de obras particulares: código em série, cliente, escopo, mão de obra,
materiais (import de fornecedor), fotos/croquis, Proposta Comercial e Relatório de Conclusão em PDF."""
import csv
import io
import secrets
from datetime import datetime, timezone
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from PIL import Image
from pydantic import BaseModel

from orcamento_pdf import build_conclusao_pdf, build_proposta_pdf

STATUSES = ["rascunho", "enviado", "aprovado", "em_execucao", "concluido", "cancelado"]
STATUS_LABELS = {
    "rascunho": "Rascunho", "enviado": "Enviado ao cliente", "aprovado": "Aprovado", "em_execucao": "Em execução",
    "concluido": "Concluído", "cancelado": "Cancelado",
}
PHOTO_KINDS = {"servico", "antes", "depois", "croqui"}

DEFAULT_TEXTS = {
    "labor_notes": "ONDE ESTÃO INCLUSOS TODOS OS CUSTOS DIRETOS E INDIRETOS COM MÃO DE OBRA, IMPOSTOS, TAXAS E CONTRIBUIÇÕES "
                   "PERTINENTES E ENCARGOS SOCIAIS, SEGUROS E TODOS DEMAIS CUSTOS NECESSÁRIOS À EXECUÇÃO DOS SERVIÇOS.",
    "material_notes": "FORNECIMENTO PELA CONTRATADA. LISTA DE MATERIAIS SERÁ ENCAMINHADA COM O ACEITE DO ORÇAMENTO.",
    "payment_terms": "30% DE ENTRADA E MOBILIZAÇÃO\nDEMAIS PARCELAS CONFORME FECHAMENTO DE MEDIÇÃO DE MATERIAIS E MÃO DE OBRA APLICADAS",
    "validity_text": "Esta proposta terá validade de 30 (trinta) dias, contados a partir da data de entrega da proposta.",
    "mobilization_text": "O cronograma de mobilização será realizado de imediato, conforme acertado na reunião de planejamento dos trabalhos.",
    "delivery_text": "DE IMEDIATO CONFORME PLANEJAMENTO EM CONJUNTO COM O CLIENTE.",
    "taxes_text": "Todos inclusos na composição do preço ofertado (contribuições municipais, estaduais e federais).",
    "intro_text": "Em atenção à gentil e cordial solicitação, temos a grata satisfação de enviar para vossa apreciação e consideração "
                  "nossa proposta técnica e comercial para atendê-los:",
}
DEFAULT_COMPANY = {
    "name": "HAYAH ENERGY", "tagline": "Engenharia e Serviços Elétricos",
    "legal_name": "HAYAH ENERGY", "cnpj": "",
    "address": "", "city": "Cuiabá/MT",
    "phone": "", "email": "",
    "responsible_name": "", "responsible_title": "HAYAH ENERGY",
    "responsible_phone": "", "footer": "",
    "logo_storage_path": "", "logo_content_type": "",
}


def _f(v) -> float:
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


class ClientInfo(BaseModel):
    name: str = ""
    contact: str = ""
    phone: str = ""
    email: str = ""
    document: str = ""
    address: str = ""
    city: str = ""


class LaborItem(BaseModel):
    id: Optional[str] = None
    description: str = ""
    unit: str = "SV"
    qty: float = 1
    unit_price: float = 0


class MaterialItem(BaseModel):
    id: Optional[str] = None
    code: str = ""
    description: str = ""
    unit: str = "UN"
    qty: float = 1
    unit_cost: float = 0
    supplier: str = ""


class OrcamentoBody(BaseModel):
    client: Optional[ClientInfo] = None
    objeto: Optional[str] = None
    location: Optional[str] = None
    base_date: Optional[str] = None
    validity_days: Optional[int] = None
    scope_items: Optional[List[str]] = None
    labor_items: Optional[List[LaborItem]] = None
    material_items: Optional[List[MaterialItem]] = None
    material_markup_pct: Optional[float] = None
    labor_notes: Optional[str] = None
    material_notes: Optional[str] = None
    payment_terms: Optional[str] = None
    validity_text: Optional[str] = None
    mobilization_text: Optional[str] = None
    delivery_text: Optional[str] = None
    taxes_text: Optional[str] = None
    intro_text: Optional[str] = None
    include_labor_table: Optional[bool] = None
    include_material_table: Optional[bool] = None
    completion_notes: Optional[str] = None
    executed_by: Optional[str] = None
    started_at: Optional[str] = None
    finished_at: Optional[str] = None
    notes: Optional[str] = None


class StatusBody(BaseModel):
    status: str


class CompanyBody(BaseModel):
    name: str = ""
    tagline: str = ""
    legal_name: str = ""
    cnpj: str = ""
    address: str = ""
    city: str = ""
    phone: str = ""
    email: str = ""
    responsible_name: str = ""
    responsible_title: str = ""
    responsible_phone: str = ""
    footer: str = ""


def compute_totals(doc: dict) -> dict:
    markup = _f(doc.get("material_markup_pct"))
    labor = []
    for it in doc.get("labor_items") or []:
        total = round(_f(it.get("qty")) * _f(it.get("unit_price")), 2)
        labor.append({**it, "qty": _f(it.get("qty")), "unit_price": _f(it.get("unit_price")), "total": total})
    mats = []
    for it in doc.get("material_items") or []:
        unit_price = round(_f(it.get("unit_cost")) * (1 + markup / 100.0), 2)
        total = round(_f(it.get("qty")) * unit_price, 2)
        mats.append({**it, "qty": _f(it.get("qty")), "unit_cost": _f(it.get("unit_cost")), "unit_price": unit_price, "total": total})
    labor_total = round(sum(i["total"] for i in labor), 2)
    material_total = round(sum(i["total"] for i in mats), 2)
    material_cost_total = round(sum(i["qty"] * i["unit_cost"] for i in mats), 2)
    return {
        "labor_items": labor, "material_items": mats,
        "labor_total": labor_total, "material_total": material_total, "material_cost_total": material_cost_total,
        "grand_total": round(labor_total + material_total, 2),
    }


def _parse_material_file(content: bytes, filename: str) -> List[dict]:
    """Planilha de fornecedor/distribuidor (xlsx/csv): detecta cabeçalho por palavras-chave."""
    keys = {
        "code": ["cod", "código", "codigo", "ref", "sku", "item mat"],
        "description": ["descri", "material", "produto", "item", "nome"],
        "qty": ["qtd", "quant", "qty", "qte"],
        "unit": ["unid", "un.", "un", "um", "unit"],
        "cost": ["preço unit", "preco unit", "valor unit", "vlr unit", "unit price", "preço", "preco", "valor", "custo", "price", "vlr"],
    }

    def to_float(v) -> float:
        s = str(v or "").strip().replace("R$", "").strip()
        if not s:
            return 0.0
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
        elif "," in s:
            s = s.replace(",", ".")
        try:
            return float(s)
        except Exception:
            return 0.0

    def classify(h: str):
        h = str(h or "").strip().lower()
        if not h:
            return None
        if any(k in h for k in ["total"]):
            return None
        if any(k in h for k in keys["cost"]):
            return "cost"
        if any(k in h for k in keys["qty"]):
            return "qty"
        if any(k in h for k in keys["code"]):
            return "code"
        if any(k in h for k in keys["unit"]) and len(h) <= 12:
            return "unit"
        if any(k in h for k in keys["description"]):
            return "description"
        return None

    rows: List[list] = []
    fname = (filename or "").lower()
    if fname.endswith(".csv") or fname.endswith(".txt"):
        text = content.decode("utf-8-sig", errors="ignore")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t") if text.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Arquivo inválido — envie uma planilha .xlsx ou .csv")
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                rows.append(list(r))
    header_idx, fields = None, None
    for i, r in enumerate(rows[:60]):
        cls = [classify(c) for c in r]
        if "description" in cls and ("qty" in cls or "cost" in cls):
            header_idx, fields = i, cls
            break
    if header_idx is None:
        raise HTTPException(status_code=400, detail="Não encontrei cabeçalho com Descrição + Quantidade/Preço na planilha")
    col = {}
    for i, f in enumerate(fields):
        if f and f not in col:
            col[f] = i
    out = []
    for r in rows[header_idx + 1:]:
        if not r or not any(c not in (None, "") for c in r):
            continue
        desc = str(r[col["description"]] if col["description"] < len(r) else "").strip()
        if not desc or desc.lower() in ("total", "totais"):
            continue
        qty = to_float(r[col["qty"]]) if "qty" in col and col["qty"] < len(r) else 1
        out.append({
            "id": secrets.token_urlsafe(6), "code": str(r[col["code"]] or "").strip() if "code" in col and col["code"] < len(r) and r[col["code"]] is not None else "",
            "description": desc, "unit": (str(r[col["unit"]] or "UN").strip().upper() if "unit" in col and col["unit"] < len(r) else "UN")[:8],
            "qty": qty if qty > 0 else 1, "unit_cost": to_float(r[col["cost"]]) if "cost" in col and col["cost"] < len(r) else 0.0,
        })
    if not out:
        raise HTTPException(status_code=400, detail="Nenhum item válido encontrado na planilha")
    return out


def build_router(db, get_current_user, safe_object_id, storage_put, storage_get, app_name: str, today_fn):
    router = APIRouter(prefix="/api")

    async def _company(user_id) -> dict:
        doc = await db.app_settings.find_one({"user_id": user_id, "key": "company_profile"})
        return {**DEFAULT_COMPANY, **((doc or {}).get("value") or {})}

    def _public(doc: dict) -> dict:
        d = {k: v for k, v in doc.items() if k not in ("_id", "user_id")}
        d["id"] = str(doc["_id"])
        d.update(compute_totals(doc))
        d["status_label"] = STATUS_LABELS.get(d.get("status", ""), d.get("status", ""))
        d["photos"] = [{**p, "url": f"/api/orcamentos/{d['id']}/photos/{p.get('id')}/file"} for p in (doc.get("photos") or [])]
        return d

    async def _get(oid: str, user) -> dict:
        doc = await db.orcamentos.find_one({"_id": safe_object_id(oid), "user_id": user["_id"]})
        if not doc:
            raise HTTPException(status_code=404, detail="Orçamento não encontrado")
        return doc

    async def _next_code(user_id) -> tuple:
        year = int(today_fn()[:4])
        last = await db.orcamentos.find_one({"user_id": user_id, "year": year}, sort=[("seq", -1)])
        seq = int((last or {}).get("seq") or 0) + 1
        return f"{seq:05d}/{year}", seq, year

    # ---------------- empresa (cabeçalho dos PDFs)
    @router.get("/orcamentos/company")
    async def get_company(current_user: dict = Depends(get_current_user)):
        c = await _company(current_user["_id"])
        c["logo_url"] = "/api/orcamentos/company/logo/file" if c.get("logo_storage_path") else ""
        return c

    @router.put("/orcamentos/company")
    async def put_company(body: CompanyBody, current_user: dict = Depends(get_current_user)):
        cur = await _company(current_user["_id"])
        value = {**cur, **body.model_dump()}
        await db.app_settings.update_one({"user_id": current_user["_id"], "key": "company_profile"},
                                         {"$set": {"value": value, "updated_at": _now()}}, upsert=True)
        return {"status": "ok", "company": value}

    @router.post("/orcamentos/company/logo")
    async def upload_logo(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
        content = await file.read()
        if not content or len(content) > 5 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Logo inválida ou maior que 5MB")
        ct = file.content_type or "image/png"
        path = f"{app_name}/orcamentos/{current_user['_id']}/logo-{secrets.token_urlsafe(4)}.{ 'png' if 'png' in ct else 'jpg'}"
        res = await storage_put(path, content, ct)
        cur = await _company(current_user["_id"])
        value = {**cur, "logo_storage_path": res.get("path", path), "logo_content_type": ct}
        await db.app_settings.update_one({"user_id": current_user["_id"], "key": "company_profile"},
                                         {"$set": {"value": value, "updated_at": _now()}}, upsert=True)
        return {"status": "ok", "logo_url": "/api/orcamentos/company/logo/file"}

    @router.get("/orcamentos/company/logo/file")
    async def logo_file(current_user: dict = Depends(get_current_user)):
        c = await _company(current_user["_id"])
        if not c.get("logo_storage_path"):
            raise HTTPException(status_code=404, detail="Sem logo")
        data, ct = await storage_get(c["logo_storage_path"])
        return Response(content=data, media_type=c.get("logo_content_type") or ct, headers={"Cache-Control": "private, max-age=3600"})

    # ---------------- CRUD
    @router.get("/orcamentos")
    async def list_orcamentos(q: str = "", status: str = "", current_user: dict = Depends(get_current_user)):
        query = {"user_id": current_user["_id"]}
        if status:
            query["status"] = status
        docs = await db.orcamentos.find(query).sort([("year", -1), ("seq", -1)]).to_list(2000)
        out = []
        ql = q.strip().lower()
        for d in docs:
            hay = f"{d.get('code','')} {(d.get('client') or {}).get('name','')} {d.get('objeto','')} {d.get('location','')}".lower()
            if ql and ql not in hay:
                continue
            t = compute_totals(d)
            out.append({
                "id": str(d["_id"]), "code": d.get("code"), "status": d.get("status"), "status_label": STATUS_LABELS.get(d.get("status", ""), ""),
                "client_name": (d.get("client") or {}).get("name", ""), "objeto": d.get("objeto", ""), "location": d.get("location", ""),
                "base_date": d.get("base_date", ""), "created_at": d.get("created_at"), "updated_at": d.get("updated_at"),
                "labor_total": t["labor_total"], "material_total": t["material_total"], "grand_total": t["grand_total"],
                "photos_count": len(d.get("photos") or []), "labor_count": len(d.get("labor_items") or []), "material_count": len(d.get("material_items") or []),
            })
        return out

    @router.post("/orcamentos")
    async def create_orcamento(body: OrcamentoBody, current_user: dict = Depends(get_current_user)):
        code, seq, year = await _next_code(current_user["_id"])
        doc = {
            "user_id": current_user["_id"], "code": code, "seq": seq, "year": year, "status": "rascunho",
            "created_at": _now(), "updated_at": _now(), "created_by": current_user.get("name", "") or current_user.get("email", ""),
            "base_date": body.base_date or today_fn(), "validity_days": body.validity_days or 30,
            "client": (body.client or ClientInfo()).model_dump(), "objeto": body.objeto or "", "location": body.location or "",
            "scope_items": body.scope_items or [], "labor_items": [], "material_items": [], "material_markup_pct": _f(body.material_markup_pct),
            "include_labor_table": False if body.include_labor_table is None else body.include_labor_table,
            "include_material_table": False if body.include_material_table is None else body.include_material_table,
            "photos": [], "notes": body.notes or "", "completion_notes": "", "executed_by": "", "started_at": "", "finished_at": "",
            **DEFAULT_TEXTS,
        }
        for k in DEFAULT_TEXTS:
            v = getattr(body, k, None)
            if v is not None:
                doc[k] = v
        r = await db.orcamentos.insert_one(doc)
        doc["_id"] = r.inserted_id
        return _public(doc)

    @router.get("/orcamentos/{oid}")
    async def get_orcamento(oid: str, current_user: dict = Depends(get_current_user)):
        return _public(await _get(oid, current_user))

    @router.put("/orcamentos/{oid}")
    async def update_orcamento(oid: str, body: OrcamentoBody, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        upd = {}
        data = body.model_dump(exclude_none=True)
        for k, v in data.items():
            if k == "client":
                upd["client"] = {**(doc.get("client") or {}), **v}
            elif k in ("labor_items", "material_items"):
                items = []
                for it in v:
                    it = dict(it)
                    it["id"] = it.get("id") or secrets.token_urlsafe(6)
                    items.append(it)
                upd[k] = items
            elif k == "scope_items":
                upd[k] = [s.strip() for s in v if str(s).strip()]
            else:
                upd[k] = v
        if not upd:
            raise HTTPException(status_code=400, detail="Nada para atualizar")
        upd["updated_at"] = _now()
        await db.orcamentos.update_one({"_id": doc["_id"]}, {"$set": upd})
        return _public({**doc, **upd})

    @router.delete("/orcamentos/{oid}")
    async def delete_orcamento(oid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        await db.orcamentos.delete_one({"_id": doc["_id"]})
        return {"status": "deleted"}

    @router.post("/orcamentos/{oid}/status")
    async def set_status(oid: str, body: StatusBody, current_user: dict = Depends(get_current_user)):
        if body.status not in STATUSES:
            raise HTTPException(status_code=400, detail=f"Status inválido. Válidos: {STATUSES}")
        doc = await _get(oid, current_user)
        upd = {"status": body.status, "updated_at": _now()}
        if body.status == "em_execucao" and not doc.get("started_at"):
            upd["started_at"] = today_fn()
        if body.status == "concluido" and not doc.get("finished_at"):
            upd["finished_at"] = today_fn()
        hist = {"status": body.status, "at": _now(), "by": current_user.get("name", "") or current_user.get("email", "")}
        await db.orcamentos.update_one({"_id": doc["_id"]}, {"$set": upd, "$push": {"status_history": hist}})
        return {"status": "ok", **upd}

    @router.post("/orcamentos/{oid}/duplicate")
    async def duplicate(oid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        code, seq, year = await _next_code(current_user["_id"])
        new = {k: v for k, v in doc.items() if k not in ("_id", "status_history", "photos")}
        new.update({"code": code, "seq": seq, "year": year, "status": "rascunho", "created_at": _now(), "updated_at": _now(),
                    "base_date": today_fn(), "photos": [], "completion_notes": "", "started_at": "", "finished_at": "", "source_code": doc.get("code")})
        r = await db.orcamentos.insert_one(new)
        new["_id"] = r.inserted_id
        return _public(new)

    # ---------------- materiais: import de planilha do fornecedor
    @router.post("/orcamentos/{oid}/materials/import")
    async def import_materials(oid: str, file: UploadFile = File(...), supplier: str = Form(""), replace: bool = Form(False),
                               current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        content = await file.read()
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Arquivo maior que 10MB")
        items = _parse_material_file(content, file.filename or "")
        for it in items:
            it["supplier"] = supplier.strip()
        final = items if replace else (doc.get("material_items") or []) + items
        await db.orcamentos.update_one({"_id": doc["_id"]}, {"$set": {"material_items": final, "updated_at": _now()}})
        return {"imported": len(items), "total_items": len(final), "items": compute_totals({**doc, "material_items": final})["material_items"]}

    # ---------------- fotos & croquis
    @router.post("/orcamentos/{oid}/photos")
    async def upload_photo(oid: str, file: UploadFile = File(...), kind: str = Form("servico"), caption: str = Form(""),
                           current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        if kind not in PHOTO_KINDS:
            raise HTTPException(status_code=400, detail=f"Tipo inválido. Use: {sorted(PHOTO_KINDS)}")
        content = await file.read()
        if not content or len(content) > 15 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Imagem vazia ou maior que 15MB")
        ct = file.content_type or "image/jpeg"
        pid = secrets.token_urlsafe(8)
        ext = "png" if "png" in ct else "jpg"
        path = f"{app_name}/orcamentos/{current_user['_id']}/{oid}/{pid}.{ext}"
        res = await storage_put(path, content, ct)
        photo = {"id": pid, "kind": kind, "caption": caption.strip(), "storage_path": res.get("path", path), "content_type": ct,
                 "size": len(content), "uploaded_at": _now(), "uploaded_by": current_user.get("name", "")}
        await db.orcamentos.update_one({"_id": doc["_id"]}, {"$push": {"photos": photo}, "$set": {"updated_at": _now()}})
        return {**photo, "url": f"/api/orcamentos/{oid}/photos/{pid}/file"}

    @router.get("/orcamentos/{oid}/photos/{pid}/file")
    async def photo_file(oid: str, pid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        ph = next((p for p in (doc.get("photos") or []) if p.get("id") == pid), None)
        if not ph:
            raise HTTPException(status_code=404, detail="Imagem não encontrada")
        data, ct = await storage_get(ph["storage_path"])
        return Response(content=data, media_type=ph.get("content_type") or ct, headers={"Cache-Control": "private, max-age=86400"})

    @router.patch("/orcamentos/{oid}/photos/{pid}")
    async def patch_photo(oid: str, pid: str, body: dict, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        upd = {}
        if "caption" in body:
            upd["photos.$.caption"] = str(body.get("caption") or "").strip()
        if body.get("kind") in PHOTO_KINDS:
            upd["photos.$.kind"] = body["kind"]
        if not upd:
            raise HTTPException(status_code=400, detail="Nada para atualizar")
        await db.orcamentos.update_one({"_id": doc["_id"], "photos.id": pid}, {"$set": upd})
        return {"status": "ok"}

    @router.delete("/orcamentos/{oid}/photos/{pid}")
    async def delete_photo(oid: str, pid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        await db.orcamentos.update_one({"_id": doc["_id"]}, {"$pull": {"photos": {"id": pid}}, "$set": {"updated_at": _now()}})
        return {"status": "deleted"}

    # ---------------- PDFs
    async def _assets(doc: dict, user) -> tuple:
        company = await _company(user["_id"])
        logo = None
        if company.get("logo_storage_path"):
            try:
                logo, _ = await storage_get(company["logo_storage_path"])
            except Exception:
                logo = None
        photos = []
        for p in (doc.get("photos") or []):
            try:
                data, _ = await storage_get(p["storage_path"])
                Image.open(io.BytesIO(data)).verify()
                photos.append({**p, "bytes": data})
            except Exception:
                continue
        return company, logo, photos

    def _fname(prefix: str, doc: dict) -> str:
        code = (doc.get("code") or "").replace("/", "-")
        return f"{prefix}_{code}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"

    @router.get("/orcamentos/{oid}/pdf/proposta")
    async def pdf_proposta(oid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        company, logo, photos = await _assets(doc, current_user)
        pdf = build_proposta_pdf(_public(doc), company, logo, photos)
        return Response(content=pdf, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{_fname("proposta_comercial", doc)}"'})

    @router.get("/orcamentos/{oid}/pdf/conclusao")
    async def pdf_conclusao(oid: str, current_user: dict = Depends(get_current_user)):
        doc = await _get(oid, current_user)
        company, logo, photos = await _assets(doc, current_user)
        pdf = build_conclusao_pdf(_public(doc), company, logo, photos)
        return Response(content=pdf, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{_fname("relatorio_conclusao", doc)}"'})

    return router
