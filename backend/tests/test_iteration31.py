"""
Iteration 31 - Relatório de Movimentação de Materiais (ledger) + Auditoria.

Covers:
- GET /api/projects/{id}/materials/ledger (structure, summary, chronological events, coherence)
- POST /api/projects/{id}/materials/add-manual → event type='avulso'
- PATCH /api/projects/{id}/materials/{item_id} → event type='edit'; no-op no event; 400/404
- POST /api/projects/{id}/materials-import → event type='import' + imported_at
- POST /api/materials/request (parcial/total) → derived requisicao events
- POST .../requisitions/{req_id}/fulfillment → separado/entregue events (only if qty>0)
- POST .../requisitions/{req_id}/resend → event type='reenvio' (general)
- POST .../materials/{req_id}/response → event type='resposta_almox' (general)
- POST .../requisitions/{req_id}/confirm-receipt → event type='recebimento' (general)
- POST .../materials/returns → event type='devolucao' by item
- GET .../ledger/pdf → 200 application/pdf, contains expected phrases (pypdf)
- GET .../ledger/xlsx → 200 with Resumo/Eventos/Obra sheets (openpyxl)
- Empty project / 404 / cross-user cases
"""
import io
import json
import os
import re
import time
from datetime import date, timedelta

import pytest
import requests
from openpyxl import load_workbook
from pypdf import PdfReader


def _load_backend_url():
    v = os.environ.get("REACT_APP_BACKEND_URL", "").strip()
    if v:
        return v.rstrip("/")
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("REACT_APP_BACKEND_URL="):
                return ln.split("=", 1)[1].strip().rstrip("/")
    return ""


BASE_URL = _load_backend_url()
API = f"{BASE_URL}/api"
ADMIN_EMAIL = "lucasc.cuiaba02@gmail.com"
ADMIN_PASS = "Hayah$152610"


@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    assert r.status_code == 200, f"login: {r.status_code} {r.text}"
    tok = r.json().get("token") or r.json().get("access_token")
    if tok:
        s.headers.update({"Authorization": f"Bearer {tok}"})
    return s


@pytest.fixture(scope="module")
def project(session):
    suffix = str(int(time.time()))
    payload = {
        "name": f"TEST_LEDGER_{suffix}",
        "project_code": f"TSTLED-{suffix}",
        "start_date": date.today().isoformat(),
        "end_date": (date.today() + timedelta(days=30)).isoformat(),
        "programmed_value": 100000,
        "daily_goal": 5000,
        "team_name": "Equipe Ledger",
    }
    r = session.post(f"{API}/projects", json=payload, timeout=30)
    assert r.status_code in (200, 201), f"create project: {r.status_code} {r.text}"
    pj = r.json()
    pid = pj.get("id") or pj.get("_id")
    ctx = {"id": pid, "code": payload["project_code"]}
    yield ctx
    try:
        session.delete(f"{API}/projects/{pid}", timeout=30)
    except Exception:
        pass


# ----- 1. add-manual generates 'avulso' event ----------------------------
def test_add_manual_creates_avulso_event(session, project):
    r = session.post(
        f"{API}/projects/{project['id']}/materials/add-manual",
        json={"item": "Poste 11m", "code": "P11", "quantity": 20, "unit": "un", "unit_cost": 500.0},
        timeout=30,
    )
    assert r.status_code in (200, 201), r.text
    mat_poste = r.json()["item"]["id"]
    project["mat_poste"] = mat_poste

    r2 = session.post(
        f"{API}/projects/{project['id']}/materials/add-manual",
        json={"item": "Cabo 35mm", "code": "C35", "quantity": 500, "unit": "m", "unit_cost": 12.0},
        timeout=30,
    )
    project["mat_cabo"] = r2.json()["item"]["id"]

    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    avulso_events = [e for e in led["events"] if e["type"] == "avulso"]
    assert len(avulso_events) == 2
    poste_ev = [e for e in avulso_events if e["material_id"] == mat_poste][0]
    assert poste_ev["qty"] == 20
    assert poste_ev["type_label"] == "Avulso adicionado"
    assert poste_ev["by"]  # not empty


# ----- 2. PATCH generates 'edit' event with old/new -----------------------
def test_patch_edit_generates_edit_event(session, project):
    pid = project["id"]
    mid = project["mat_poste"]
    # First set qty_requested to a known value
    r = session.patch(f"{API}/projects/{pid}/materials/{mid}", json={"qty_requested": 10}, timeout=30)
    assert r.status_code == 200
    # Then change 10 → 12
    r2 = session.patch(f"{API}/projects/{pid}/materials/{mid}", json={"qty_requested": 12}, timeout=30)
    assert r2.status_code == 200

    led = session.get(f"{API}/projects/{pid}/materials/ledger", timeout=30).json()
    edits = [e for e in led["events"] if e["type"] == "edit" and e["material_id"] == mid]
    assert len(edits) >= 1
    last = edits[-1]
    assert "Qtd. solicitada" in last["detail"]
    assert "→" in last["detail"]
    assert "12" in last["detail"]


def test_patch_no_change_does_not_log(session, project):
    pid = project["id"]
    mid = project["mat_cabo"]
    r = session.patch(f"{API}/projects/{pid}/materials/{mid}", json={"qty_requested": 500}, timeout=30)
    assert r.status_code == 200
    before_led = session.get(f"{API}/projects/{pid}/materials/ledger", timeout=30).json()
    edits_before = [e for e in before_led["events"] if e["type"] == "edit" and e["material_id"] == mid]

    # Second PATCH with same value → no new edit event
    r2 = session.patch(f"{API}/projects/{pid}/materials/{mid}", json={"qty_requested": 500}, timeout=30)
    assert r2.status_code == 200
    after_led = session.get(f"{API}/projects/{pid}/materials/ledger", timeout=30).json()
    edits_after = [e for e in after_led["events"] if e["type"] == "edit" and e["material_id"] == mid]
    assert len(edits_after) == len(edits_before)


def test_patch_404_nonexistent(session, project):
    r = session.patch(f"{API}/projects/{project['id']}/materials/nope-xxx", json={"qty_requested": 1}, timeout=30)
    assert r.status_code == 404


def test_patch_400_no_valid_field(session, project):
    r = session.patch(
        f"{API}/projects/{project['id']}/materials/{project['mat_poste']}",
        json={"unknown_field": "x"}, timeout=30,
    )
    assert r.status_code == 400


# ----- 3. requisicao events on partial and total ------------------------
def test_partial_requisition_creates_requisicao_events(session, project):
    items = json.dumps([{"id": project["mat_poste"], "qty": 5}, {"id": project["mat_cabo"], "qty": 100}])
    r = session.post(
        f"{API}/materials/request",
        data={"project_id": project["id"], "signature_name": "TesterX",
              "requisition_type": "parcial", "items": items},
        timeout=60,
    )
    assert r.status_code == 200, r.text
    project["_req_partial"] = r.json()["requisition_id"]

    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    reqs = [e for e in led["events"] if e["type"] == "requisicao"]
    # Two items in the partial → 2 requisicao events (may include prior none)
    partial_evs = [e for e in reqs if "Parcial" in (e["detail"] or "")]
    assert len(partial_evs) == 2
    for e in partial_evs:
        assert "assinado por" in e["detail"]
        assert e["ref"].startswith("Doc ")
        assert e["qty"] in (5, 100)


def test_total_requisition_creates_requisicao_events_from_snapshot(session, project):
    r = session.post(
        f"{API}/materials/request",
        data={"project_id": project["id"], "signature_name": "TotalTester"},
        timeout=60,
    )
    assert r.status_code == 200
    project["_req_total"] = r.json()["requisition_id"]

    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    total_evs = [e for e in led["events"] if e["type"] == "requisicao" and "Total" in (e["detail"] or "")]
    # snapshot has 2 items
    assert len(total_evs) == 2


# ----- 4. resend produces reenvio event without material_id --------------
def test_resend_creates_reenvio_general_event(session, project):
    r = session.post(f"{API}/projects/{project['id']}/requisitions/{project['_req_partial']}/resend", timeout=60)
    assert r.status_code == 200
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    reenvios = [e for e in led["events"] if e["type"] == "reenvio"]
    assert len(reenvios) >= 1
    assert all(e["material_id"] is None for e in reenvios)


# ----- 5. fulfillment produces separado/entregue events -----------------
def test_fulfillment_creates_separado_entregue_events(session, project):
    rid = project["_req_partial"]
    body = {"responsible": "AlmoxLedger", "items": [
        {"material_id": project["mat_poste"], "qty_separated": 5, "qty_delivered": 4},
        {"material_id": project["mat_cabo"], "qty_separated": 100, "qty_delivered": 0},  # entregue skipped
    ]}
    r = session.post(f"{API}/projects/{project['id']}/requisitions/{rid}/fulfillment", json=body, timeout=30)
    assert r.status_code == 200, r.text
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    sep = [e for e in led["events"] if e["type"] == "separado"]
    ent = [e for e in led["events"] if e["type"] == "entregue"]
    assert len(sep) == 2  # both items had qty_separated > 0
    assert len(ent) == 1  # only poste had qty_delivered > 0
    assert ent[0]["by"] == "AlmoxLedger"
    assert ent[0]["material_id"] == project["mat_poste"]


# ----- 6. warehouse response produces resposta_almox general event -------
def test_warehouse_response_event(session, project):
    rid = project["_req_partial"]
    r = session.post(
        f"{API}/projects/{project['id']}/materials/{rid}/response",
        json={"status": "separado", "responsible": "RespX", "notes": "ok"},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    resp = [e for e in led["events"] if e["type"] == "resposta_almox"]
    assert len(resp) >= 1
    assert all(e["material_id"] is None for e in resp)
    assert "Status" in resp[-1]["detail"]


# ----- 7. confirm-receipt general event ---------------------------------
def test_confirm_receipt_event(session, project):
    rid = project["_req_partial"]
    r = session.post(
        f"{API}/projects/{project['id']}/requisitions/{rid}/confirm-receipt",
        data={"receiver_name": "ReceberX"},
        timeout=30,
    )
    # may accept form or json - allow both
    if r.status_code not in (200, 201):
        r = session.post(
            f"{API}/projects/{project['id']}/requisitions/{rid}/confirm-receipt",
            json={"receiver_name": "ReceberX"}, timeout=30,
        )
    assert r.status_code in (200, 201), r.text
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    rec = [e for e in led["events"] if e["type"] == "recebimento"]
    assert len(rec) >= 1
    assert all(e["material_id"] is None for e in rec)


# ----- 8. return produces devolucao per item -----------------------------
def test_return_creates_devolucao_events(session, project):
    r = session.post(
        f"{API}/projects/{project['id']}/materials/returns",
        json={"date": date.today().isoformat(), "returned_by": "DevX", "reason": "sobra", "notes": "n1",
              "items": [{"material_id": project["mat_poste"], "qty": 2}]},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    devs = [e for e in led["events"] if e["type"] == "devolucao"]
    assert len(devs) >= 1
    d = devs[-1]
    assert d["by"] == "DevX"
    assert d["material_id"] == project["mat_poste"]
    assert "sobra" in d["detail"]
    assert d["qty"] == 2


# ----- 9. ledger summary/items coherence & chronological ordering -------
def test_ledger_summary_coherence_and_ordering(session, project):
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    # Required top-level fields
    for key in ("project", "generated_at", "summary", "items", "events"):
        assert key in led
    pj = led["project"]
    for key in ("code", "name", "team_name", "location"):
        assert key in pj

    s = led["summary"]
    for key in ("items_count", "total_qty", "requisitioned_qty", "returned_qty", "in_field_qty",
                "separated_qty", "delivered_qty", "saldo_qty", "total_cost", "in_field_cost",
                "requisitions_count", "returns_count", "events_count", "items_fully_requisitioned"):
        assert key in s, f"missing summary key: {key}"

    # events_count matches
    assert s["events_count"] == len(led["events"])

    # coherence per item: in_field = req - ret; saldo = qty_total - in_field (min 0);
    # in_field_cost = unit_cost * in_field
    sum_req = sum(i["qty_requisitioned"] for i in led["items"])
    sum_ret = sum(i["qty_returned"] for i in led["items"])
    for it in led["items"]:
        assert abs(it["in_field"] - max(0.0, it["qty_requisitioned"] - it["qty_returned"])) < 1e-6
        assert abs(it["saldo"] - max(0.0, it["qty_total"] - it["in_field"])) < 1e-6
        assert abs(it["in_field_cost"] - round(it["unit_cost"] * it["in_field"], 2)) < 1e-2

    assert abs(s["requisitioned_qty"] - sum_req) < 1e-6
    assert abs(s["returned_qty"] - sum_ret) < 1e-6

    # chronological events
    ts_list = [e["ts"] for e in led["events"]]
    assert ts_list == sorted(ts_list)


# ----- 10. PDF endpoint --------------------------------------------------
def test_ledger_pdf(session, project):
    r = session.get(f"{API}/projects/{project['id']}/materials/ledger/pdf", timeout=60)
    assert r.status_code == 200, r.text[:200]
    assert "application/pdf" in r.headers.get("Content-Type", "")
    cd = r.headers.get("Content-Disposition", "")
    assert "attachment" in cd
    assert "movimentacao_materiais" in cd
    txt = "".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(r.content)).pages)
    assert "Relatório de Movimentação de Materiais" in txt
    assert "Histórico de acontecimentos por material" in txt


# ----- 11. XLSX endpoint -------------------------------------------------
def test_ledger_xlsx(session, project):
    r = session.get(f"{API}/projects/{project['id']}/materials/ledger/xlsx", timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) >= {"Resumo", "Eventos", "Obra"}

    ws_ev = wb["Eventos"]
    headers = [c.value for c in ws_ev[1]]
    expected = ["Data/Hora", "Tipo", "Código", "Material", "Un", "Qtd", "Responsável", "Referência", "Detalhes"]
    assert headers == expected

    # There should be at least one event row + header
    assert ws_ev.max_row >= 2

    ws_r = wb["Resumo"]
    # There should be a TOTAIS row somewhere in Resumo
    found_totais = False
    for row in ws_r.iter_rows(values_only=True):
        if row and row[0] == "TOTAIS":
            found_totais = True
            break
    assert found_totais, "TOTAIS row missing in Resumo sheet"


# ----- 12. Empty ledger + 404 --------------------------------------------
def test_ledger_empty_project(session):
    suffix = str(int(time.time()))
    r = session.post(f"{API}/projects", json={
        "name": f"TEST_LEDGER_EMPTY_{suffix}", "project_code": f"TLEEM-{suffix}",
        "start_date": date.today().isoformat(), "end_date": (date.today() + timedelta(days=10)).isoformat(),
        "programmed_value": 1, "daily_goal": 1,
    }, timeout=30)
    assert r.status_code in (200, 201)
    pid = r.json().get("id") or r.json().get("_id")
    try:
        led = session.get(f"{API}/projects/{pid}/materials/ledger", timeout=30).json()
        assert led["items"] == []
        s = led["summary"]
        for k in ("total_qty", "requisitioned_qty", "returned_qty", "in_field_qty",
                  "saldo_qty", "total_cost", "events_count", "requisitions_count"):
            assert s[k] == 0, f"{k} should be 0"
        assert s["items_count"] == 0
    finally:
        session.delete(f"{API}/projects/{pid}", timeout=30)


def test_ledger_404_invalid_project(session):
    r = session.get(f"{API}/projects/000000000000000000000000/materials/ledger", timeout=30)
    assert r.status_code == 404


# ----- 13. Import creates 'import' events -------------------------------
def test_import_generates_import_events(session, project):
    csv = (
        "item;quantidade;unidade;preco\n"
        "TESTIMP-Cimento 50kg;3;sc;35\n"
        "TESTIMP-Areia lavada;2;m3;120\n"
    )
    files = {"file": ("t.csv", csv.encode("utf-8"), "text/csv")}
    r = session.post(f"{API}/projects/{project['id']}/materials-import", files=files, timeout=60)
    assert r.status_code in (200, 201), r.text
    led = session.get(f"{API}/projects/{project['id']}/materials/ledger", timeout=30).json()
    imp = [e for e in led["events"] if e["type"] == "import"]
    # 1 general import event + 1 per imported item (2)
    assert len(imp) >= 3
    general = [e for e in imp if e["material_id"] is None]
    assert len(general) >= 1
    # imported items should also carry imported_at → per-item import events
    per_item = [e for e in imp if e["material_id"] is not None]
    assert len(per_item) >= 2
