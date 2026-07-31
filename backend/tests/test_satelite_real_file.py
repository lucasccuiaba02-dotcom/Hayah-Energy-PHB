"""Test the smart parser against the real Satelite-atualxxMODELO.xlsx file
provided by the user. Verifies:
 1. Sheet REQ. DE MATERIAL-GERAL is preferred over DIARIO DE OBRA (via sheet_priority)
 2. Header is correctly detected on row 15 (with logo/metadata above)
 3. #N/A rows are filtered out
 4. QNTD ORÇADA is mapped to quantity (not QNTD SEPARADA/USADA/DEVOLVIDA)
Also tests a "filled" version by injecting data into the template.
"""
import io
import os
import shutil
import httpx
import pytest
from openpyxl import load_workbook

BACKEND_URL = os.environ.get("REACT_APP_BACKEND_URL") or "http://localhost:8001"
API = BACKEND_URL.rstrip("/") + ("" if BACKEND_URL.endswith("/api") else "/api")

ADMIN_EMAIL = "admin@projectmatrix.com"
ADMIN_PWD = "Admin@123"

SATELITE_PATH = os.path.join(os.path.dirname(__file__), "fixtures", "satelite.xlsx")


def _login():
    r = httpx.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PWD}, timeout=15)
    r.raise_for_status()
    return r.json()["access_token"]


def _create_project(token, name):
    import time
    r = httpx.post(f"{API}/projects", json={
        "name": name,
        "project_code": f"TEST-SAT-{name[:8]}-{int(time.time()*1000) % 10**8}",
        "programmed_value": 1000,
    }, headers={"Authorization": f"Bearer {token}"}, timeout=15)
    r.raise_for_status()
    d = r.json()
    return d.get("id") or d.get("_id")


def _delete_project(token, pid):
    httpx.delete(f"{API}/projects/{pid}", headers={"Authorization": f"Bearer {token}"}, timeout=15)


@pytest.fixture(scope="module")
def token():
    return _login()


def _make_filled_satelite() -> bytes:
    """Load the real Satelite template and inject 5 fake material rows in REQ. DE MATERIAL-GERAL
    (rows 16-20) so the parser has something to extract."""
    assert os.path.exists(SATELITE_PATH), f"Missing Satelite file at {SATELITE_PATH}"
    wb = load_workbook(SATELITE_PATH)
    ws = wb["REQ. DE MATERIAL-GERAL"]
    fills = [
        ("PADRAO", "12345", "POSTE DE CONCRETO DT 10/300", "UN", 5, 0, 0, 0),
        ("PADRAO", "23456", "CABO DE ALUMINIO 4/0AWG", "M", 250, 100, 50, 0),
        ("PADRAO", "34567", "CRUZETA MADEIRA 2,4M", "UN", 8, 0, 0, 0),
        ("PADRAO", "45678", "TRANSFORMADOR 30KVA MONOF", "UN", 1, 0, 0, 0),
        ("PADRAO", "56789", "ISOLADOR POLIMERICO 15KV", "UN", 24, 12, 0, 0),
    ]
    # Row 15 headers: A=ORÇADO EM OBRA, C=CÓDIGO, D=DESCRIÇÃO, G=QNTD ORÇADA, H=QNTD SEPARADA, I=QNTD USADA, J=QNTD DEVOLVIDA
    for i, (situ, cod, desc, un, orc, sep, usa, dev) in enumerate(fills):
        row = 16 + i
        ws.cell(row=row, column=1).value = situ
        ws.cell(row=row, column=3).value = cod
        ws.cell(row=row, column=4).value = desc
        # UN is not in header explicitly (col E/F are blank), so we don't fill unit
        ws.cell(row=row, column=7).value = orc  # QNTD ORÇADA
        ws.cell(row=row, column=8).value = sep  # QNTD SEPARADA
        ws.cell(row=row, column=9).value = usa
        ws.cell(row=row, column=10).value = dev
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_handles_raw_satelite_template(token):
    """Raw template has #N/A everywhere in REQ. DE MATERIAL-GERAL, DIARIO DE OBRA also has fake ORÇADO
    data. Parser must NOT crash; result may be low (mostly filtered out)."""
    assert os.path.exists(SATELITE_PATH)
    pid = _create_project(token, "TEST_RAW_SATELITE")
    try:
        with open(SATELITE_PATH, "rb") as f:
            content = f.read()
        files = {"file": ("satelite.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = httpx.post(
            f"{API}/projects/{pid}/materials-import",
            files=files,
            data={"replace": "true"},
            headers={"Authorization": f"Bearer {token}"},
            timeout=60,
        )
        assert r.status_code == 200, r.text
        d = r.json()
        # Should not fail; items may be 0 (all #N/A filtered) or DIARIO-derived
        # We only require no crash + no #N/A rows leaking
        for it in d.get("items", []):
            assert not (it.get("item") or "").startswith("#"), f"Leaked #N/A row: {it}"
    finally:
        _delete_project(token, pid)


def test_parser_prefers_req_material_over_diario(token):
    """When Satelite is FILLED with data in REQ. DE MATERIAL-GERAL, parser must prefer it over DIARIO DE OBRA."""
    pid = _create_project(token, "TEST_FILLED_SATELITE")
    try:
        xlsx = _make_filled_satelite()
        files = {"file": ("satelite_filled.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = httpx.post(
            f"{API}/projects/{pid}/materials-import",
            files=files,
            data={"replace": "true"},
            headers={"Authorization": f"Bearer {token}"},
            timeout=60,
        )
        assert r.status_code == 200, r.text
        d = r.json()
        items = d.get("items", [])
        descriptions = {i["item"] for i in items}
        assert "POSTE DE CONCRETO DT 10/300" in descriptions, f"Missing expected material: {list(descriptions)[:10]}"
        assert "TRANSFORMADOR 30KVA MONOF" in descriptions
        # QNTD ORÇADA must be picked (not SEPARADA/USADA/DEVOLVIDA)
        cabo = next((i for i in items if i["item"] == "CABO DE ALUMINIO 4/0AWG"), None)
        assert cabo is not None
        assert cabo["quantity"] == 250, f"Expected QNTD ORÇADA=250, got {cabo}"
        # At least our 5 items (DIARIO may also contribute, but REQ should win with sheet_priority=1000)
        assert len(items) >= 5
    finally:
        _delete_project(token, pid)
