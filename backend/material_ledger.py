"""Set/2026 · Relatório de Movimentação de Materiais por obra (linha do tempo + resumo, PDF e Excel)."""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

EVENT_LABELS = {
    "import": "Lista importada",
    "avulso": "Avulso adicionado",
    "edit": "Edição manual",
    "requisicao": "Requisitado",
    "reenvio": "Reenvio da requisição",
    "separado": "Separado (almoxarifado)",
    "entregue": "Entregue (almoxarifado)",
    "resposta_almox": "Resposta do almoxarifado",
    "recebimento": "Recebimento confirmado em obra",
    "devolucao": "Devolvido ao almoxarifado",
}
FIELD_LABELS = {
    "qty_requested": "Qtd. solicitada", "qty_separated": "Separado", "qty_written_off": "Baixado",
    "qty_delivered": "Entregue", "qty_applied": "Aplicado", "unit": "Unidade", "warehouse_status": "Status almox.",
    "code": "Código", "item": "Descrição", "notes": "Observações",
}


def _f(v) -> float:
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _fmt_qty(v) -> str:
    v = _f(v)
    return f"{int(v)}" if v == int(v) else f"{v:.2f}".rstrip("0").rstrip(".")


def _fmt_brl(v) -> str:
    return f"R$ {_f(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_ts(ts: str) -> str:
    if not ts:
        return "—"
    try:
        d = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return d.strftime("%d/%m/%Y %H:%M")
    except Exception:
        s = str(ts)
        return f"{s[8:10]}/{s[5:7]}/{s[0:4]}" if len(s) >= 10 and s[4] == "-" else s


def _ev(ts, etype, material=None, qty=None, by="", ref="", detail=""):
    m = material or {}
    return {
        "ts": ts or "", "ts_label": _fmt_ts(ts), "type": etype, "type_label": EVENT_LABELS.get(etype, etype),
        "material_id": m.get("id"), "code": m.get("code", "") or "", "item": m.get("item", "") or "",
        "unit": (m.get("unit") or "UN").upper(), "qty": None if qty is None else round(_f(qty), 3),
        "by": by or "", "ref": ref or "", "detail": detail or "",
    }


def build_material_ledger(project: dict) -> dict:
    mats = project.get("materials_list") or []
    by_id = {m.get("id"): m for m in mats}
    by_name = {(m.get("item") or "").strip().lower(): m for m in mats}

    def resolve(mid=None, name=None):
        if mid and mid in by_id:
            return by_id[mid]
        return by_name.get((name or "").strip().lower())

    events = []
    # 1) itens da lista (avulsos legados sem auditoria / importados com carimbo)
    audited_avulso = {e.get("material_id") for e in (project.get("material_events") or []) if e.get("type") == "avulso"}
    for m in mats:
        if m.get("source") == "avulso" and m.get("id") not in audited_avulso and m.get("added_at"):
            events.append(_ev(m["added_at"], "avulso", m, m.get("quantity"), m.get("added_by", ""), "", m.get("notes", "")))
        elif m.get("imported_at"):
            events.append(_ev(m["imported_at"], "import", m, m.get("quantity"), m.get("imported_by", ""), "", "Item da lista orçada"))
    # 2) auditoria (edições manuais, avulsos, importações)
    for e in (project.get("material_events") or []):
        m = by_id.get(e.get("material_id")) or {"id": e.get("material_id"), "code": e.get("code", ""), "item": e.get("item", ""), "unit": e.get("unit", "UN")}
        detail = e.get("detail", "")
        if e.get("type") == "edit":
            def _v(x):
                try:
                    return _fmt_qty(float(x))
                except Exception:
                    return x
            detail = f"{FIELD_LABELS.get(e.get('field'), e.get('field'))}: {_v(e.get('old'))} → {_v(e.get('new'))}"
        events.append(_ev(e.get("ts"), e.get("type", "edit"), m if e.get("material_id") else None, e.get("qty"), e.get("by", ""), e.get("ref", ""), detail))
    # 3) requisições
    for r in (project.get("material_requisitions") or []):
        ts = r.get("created_at")
        ref = f"Doc {r.get('id', '')}"
        if r.get("is_resend"):
            events.append(_ev(ts, "reenvio", None, None, r.get("created_by", ""), ref, f"Reenvio da requisição {r.get('original_req_id', '')} · e-mail {r.get('email_status', '')}"))
            continue
        partial = r.get("requisition_type") == "parcial"
        kind = f"Parcial nº {r.get('seq') or '?'}" if partial else "Total"
        base_detail = f"{kind} · assinado por {r.get('signature_name') or '—'} · e-mail {r.get('email_status', '')}"
        if partial:
            for it in (r.get("items") or []):
                m = resolve(it.get("material_id"), it.get("item"))
                events.append(_ev(ts, "requisicao", m or {"id": it.get("material_id"), "code": it.get("code", ""), "item": it.get("item", ""), "unit": it.get("unit", "UN")},
                                  it.get("qty"), r.get("created_by", ""), ref, base_detail))
        else:
            for s in (r.get("materials_snapshot") or []):
                m = resolve(s.get("material_id"), s.get("item"))
                events.append(_ev(ts, "requisicao", m or {"id": s.get("material_id"), "code": "", "item": s.get("item", ""), "unit": s.get("unit", "UN")},
                                  s.get("qty_requested"), r.get("created_by", ""), ref, base_detail))
        ful = r.get("fulfillment") or {}
        for f in (ful.get("items") or []):
            m = by_id.get(f.get("material_id"))
            who = ful.get("responsible") or ful.get("updated_by", "")
            if _f(f.get("qty_separated")) > 0:
                events.append(_ev(ful.get("updated_at"), "separado", m, f.get("qty_separated"), who, ref, f"{kind}" + (f" · {ful.get('notes')}" if ful.get("notes") else "")))
            if _f(f.get("qty_delivered")) > 0:
                events.append(_ev(ful.get("updated_at"), "entregue", m, f.get("qty_delivered"), who, ref, f"{kind}" + (f" · {ful.get('notes')}" if ful.get("notes") else "")))
        wr = r.get("warehouse_response")
        if wr:
            d = f"Status: {wr.get('status', '')}"
            if wr.get("notes"):
                d += f" · {wr['notes']}"
            if wr.get("missing_items"):
                d += f" · faltam: {wr['missing_items']}"
            events.append(_ev(wr.get("answered_at"), "resposta_almox", None, None, wr.get("responsible") or wr.get("answered_by", ""), ref, d))
        rc = r.get("receipt")
        if rc:
            d = (rc.get("notes") or "") + (" · com foto do canhoto" if rc.get("photo_url") or rc.get("photo_path") else "")
            events.append(_ev(rc.get("received_at"), "recebimento", None, None, rc.get("received_by", ""), ref, d.strip(" ·")))
    # 4) devoluções
    for d in (project.get("material_returns") or []):
        ts = d.get("created_at") or d.get("date")
        for it in (d.get("items") or []):
            m = by_id.get(it.get("material_id")) or {"id": it.get("material_id"), "code": it.get("code", ""), "item": it.get("item", ""), "unit": it.get("unit", "UN")}
            det = " · ".join([x for x in [f"Data {_fmt_ts(d.get('date'))[:10]}", d.get("reason", ""), d.get("notes", "")] if x])
            events.append(_ev(ts, "devolucao", m, it.get("qty"), d.get("returned_by") or d.get("registered_by", ""), f"Dev {d.get('id', '')}", det))

    events.sort(key=lambda e: e["ts"] or "")

    # resumo por item
    agg = {}
    for e in events:
        if not e["material_id"]:
            continue
        a = agg.setdefault(e["material_id"], {"req": 0.0, "ret": 0.0, "sep": 0.0, "ent": 0.0})
        if e["type"] == "requisicao":
            a["req"] += _f(e["qty"])
        elif e["type"] == "devolucao":
            a["ret"] += _f(e["qty"])
        elif e["type"] == "separado":
            a["sep"] += _f(e["qty"])
        elif e["type"] == "entregue":
            a["ent"] += _f(e["qty"])
    items = []
    for m in mats:
        a = agg.get(m.get("id"), {"req": 0.0, "ret": 0.0, "sep": 0.0, "ent": 0.0})
        total = _f(m.get("qty_requested", m.get("quantity", 0)))
        in_field = max(0.0, a["req"] - a["ret"])
        unit_cost = _f(m.get("unit_cost"))
        items.append({
            "id": m.get("id"), "code": m.get("code", "") or "", "item": m.get("item", "") or "", "unit": (m.get("unit") or "UN").upper(),
            "source": m.get("source", "planilha"), "qty_total": round(total, 3),
            "qty_requisitioned": round(a["req"], 3), "qty_returned": round(a["ret"], 3), "in_field": round(in_field, 3),
            "qty_separated": round(_f(m.get("qty_separated")) or a["sep"], 3), "qty_delivered": round(_f(m.get("qty_delivered")) or a["ent"], 3),
            "qty_written_off": round(_f(m.get("qty_written_off")), 3), "qty_applied": round(_f(m.get("qty_applied")), 3),
            "saldo": round(max(0.0, total - in_field), 3), "unit_cost": round(unit_cost, 2),
            "total_cost": round(_f(m.get("total_cost")) or unit_cost * total, 2), "in_field_cost": round(unit_cost * in_field, 2),
            "returned_cost": round(unit_cost * a["ret"], 2), "warehouse_status": m.get("warehouse_status", "") or "",
            "events_count": sum(1 for e in events if e["material_id"] == m.get("id")),
        })
    summary = {
        "items_count": len(items),
        "total_qty": round(sum(i["qty_total"] for i in items), 3),
        "requisitioned_qty": round(sum(i["qty_requisitioned"] for i in items), 3),
        "returned_qty": round(sum(i["qty_returned"] for i in items), 3),
        "in_field_qty": round(sum(i["in_field"] for i in items), 3),
        "separated_qty": round(sum(i["qty_separated"] for i in items), 3),
        "delivered_qty": round(sum(i["qty_delivered"] for i in items), 3),
        "saldo_qty": round(sum(i["saldo"] for i in items), 3),
        "total_cost": round(sum(i["total_cost"] for i in items), 2),
        "in_field_cost": round(sum(i["in_field_cost"] for i in items), 2),
        "returned_cost": round(sum(i["returned_cost"] for i in items), 2),
        "requisitions_count": sum(1 for r in (project.get("material_requisitions") or []) if not r.get("is_resend")),
        "returns_count": len(project.get("material_returns") or []),
        "events_count": len(events),
        "items_fully_requisitioned": sum(1 for i in items if i["saldo"] <= 0 and i["qty_total"] > 0),
    }
    return {
        "project": {
            "id": str(project.get("_id", "")), "code": project.get("project_code", "") or "", "name": project.get("name", ""),
            "team_name": project.get("team_name", "") or "", "location": project.get("location", "") or "",
            "status": project.get("execution_status") or project.get("status", "") or "",
            "start_date": project.get("start_date", ""), "end_date": project.get("end_date", ""),
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary, "items": items, "events": events,
    }


# ---------------------------------------------------------------- Excel
def build_material_ledger_xlsx(ledger: dict) -> bytes:
    wb = Workbook()
    head_font = Font(bold=True, color="E0FF00")
    head_fill = PatternFill("solid", fgColor="050505")
    title_font = Font(bold=True, size=13)
    p = ledger["project"]
    s = ledger["summary"]

    ws = wb.active
    ws.title = "Resumo"
    ws.append([f"Relatório de Movimentação de Materiais — [{p['code']}] {p['name']}"])
    ws["A1"].font = title_font
    ws.append([f"Equipe: {p['team_name'] or '—'} · Local: {p['location'] or '—'} · Gerado em {_fmt_ts(ledger['generated_at'])}"])
    ws.append([f"Itens: {s['items_count']} · Requisições: {s['requisitions_count']} · Devoluções: {s['returns_count']} · Eventos: {s['events_count']}"])
    ws.append([])
    headers = ["Código", "Descrição", "Un", "Origem", "Qtd. Obra", "Requisitado", "Devolvido", "Em Campo", "Separado", "Entregue", "Baixado", "Saldo a Requisitar",
               "Custo Unit. (R$)", "Custo Total (R$)", "Custo em Campo (R$)", "Custo Devolvido (R$)", "Status Almox.", "Nº Eventos"]
    ws.append(headers)
    hr = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for it in ledger["items"]:
        ws.append([it["code"], it["item"], it["unit"], "Avulso" if it["source"] == "avulso" else "Planilha", it["qty_total"], it["qty_requisitioned"], it["qty_returned"],
                   it["in_field"], it["qty_separated"], it["qty_delivered"], it["qty_written_off"], it["saldo"], it["unit_cost"], it["total_cost"],
                   it["in_field_cost"], it["returned_cost"], it["warehouse_status"], it["events_count"]])
    ws.append(["TOTAIS", "", "", "", s["total_qty"], s["requisitioned_qty"], s["returned_qty"], s["in_field_qty"], s["separated_qty"], s["delivered_qty"], "",
               s["saldo_qty"], "", s["total_cost"], s["in_field_cost"], s["returned_cost"], "", s["events_count"]])
    for c in range(1, len(headers) + 1):
        ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
    widths = [14, 44, 6, 10, 11, 12, 11, 11, 11, 11, 10, 14, 14, 15, 16, 16, 14, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, min_col=13, max_col=16):
        for cell in row:
            cell.number_format = '#,##0.00'
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    ws2 = wb.create_sheet("Eventos")
    ev_headers = ["Data/Hora", "Tipo", "Código", "Material", "Un", "Qtd", "Responsável", "Referência", "Detalhes"]
    ws2.append(ev_headers)
    for c in range(1, len(ev_headers) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
    for e in ledger["events"]:
        ws2.append([e["ts_label"], e["type_label"], e["code"], e["item"] or ("(obra)" if not e["material_id"] else ""), e["unit"] if e["material_id"] else "",
                    e["qty"] if e["qty"] is not None else "", e["by"], e["ref"], e["detail"]])
    for i, w in enumerate([17, 28, 14, 44, 6, 9, 24, 18, 70], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("Obra")
    for k, v in [("Código", p["code"]), ("Nome", p["name"]), ("Equipe", p["team_name"]), ("Localidade", p["location"]), ("Status", p["status"]),
                 ("Início", p["start_date"]), ("Fim", p["end_date"]), ("Gerado em", _fmt_ts(ledger["generated_at"]))]:
        ws3.append([k, v])
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Consolidado (todas as obras · período)
def _in_period(ts: str, start: str, end: str) -> bool:
    d = (ts or "")[:10]
    if not d:
        return False
    if start and d < start[:10]:
        return False
    if end and d > end[:10]:
        return False
    return True


def _period_label(start: str, end: str) -> str:
    def f(x):
        return f"{x[8:10]}/{x[5:7]}/{x[0:4]}" if x and len(x) >= 10 else ""
    if start and end:
        return f"{f(start)} a {f(end)}"
    if start:
        return f"a partir de {f(start)}"
    if end:
        return f"até {f(end)}"
    return "todo o histórico"


def build_consolidated_xlsx(ledgers: list, start: str = "", end: str = "", team_label: str = "") -> bytes:
    wb = Workbook()
    head_font = Font(bold=True, color="E0FF00")
    head_fill = PatternFill("solid", fgColor="050505")
    bold = Font(bold=True)
    per = _period_label(start, end)

    def header(ws, cols, row=None):
        ws.append(cols)
        r = row or ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        return r

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    per_obra, por_material, eventos, posicao = [], {}, [], []
    for lg in ledgers:
        p, s = lg["project"], lg["summary"]
        evs = [e for e in lg["events"] if _in_period(e["ts"], start, end)]
        req = sum(_f(e["qty"]) for e in evs if e["type"] == "requisicao")
        ret = sum(_f(e["qty"]) for e in evs if e["type"] == "devolucao")
        ent = sum(_f(e["qty"]) for e in evs if e["type"] == "entregue")
        sep = sum(_f(e["qty"]) for e in evs if e["type"] == "separado")
        per_obra.append([p["code"], p["name"], p["team_name"], p["location"], p["status"], s["items_count"], s["total_qty"], req, ret, sep, ent,
                         s["in_field_qty"], s["saldo_qty"], s["in_field_cost"], s["total_cost"], len(evs), "Sim" if evs else "Não"])
        for e in evs:
            eventos.append([e["ts_label"], p["code"], p["name"], p["team_name"], e["type_label"], e["code"], e["item"] or ("(obra)" if not e["material_id"] else ""),
                            e["unit"] if e["material_id"] else "", e["qty"] if e["qty"] is not None else "", e["by"], e["ref"], e["detail"]])
        ev_by_item = {}
        for e in evs:
            if e["material_id"]:
                a = ev_by_item.setdefault(e["material_id"], {"req": 0.0, "ret": 0.0, "ent": 0.0})
                if e["type"] == "requisicao":
                    a["req"] += _f(e["qty"])
                elif e["type"] == "devolucao":
                    a["ret"] += _f(e["qty"])
                elif e["type"] == "entregue":
                    a["ent"] += _f(e["qty"])
        for it in lg["items"]:
            a = ev_by_item.get(it["id"], {"req": 0.0, "ret": 0.0, "ent": 0.0})
            key = (it["code"] or "").strip().upper() or f"~{(it['item'] or '').strip().lower()}"
            m = por_material.setdefault(key, {"code": it["code"], "item": it["item"], "unit": it["unit"], "obras": set(), "total": 0.0, "req": 0.0, "ret": 0.0, "ent": 0.0,
                                              "field": 0.0, "saldo": 0.0, "cost": 0.0, "unit_cost": it["unit_cost"]})
            m["obras"].add(p["code"] or p["name"])
            m["total"] += it["qty_total"]; m["req"] += a["req"]; m["ret"] += a["ret"]; m["ent"] += a["ent"]
            m["field"] += it["in_field"]; m["saldo"] += it["saldo"]; m["cost"] += it["in_field_cost"]
            if not m["item"] and it["item"]:
                m["item"] = it["item"]
            posicao.append([p["code"], p["name"], p["team_name"], it["code"], it["item"], it["unit"], "Avulso" if it["source"] == "avulso" else "Planilha", it["qty_total"],
                            it["qty_requisitioned"], it["qty_returned"], it["in_field"], it["qty_separated"], it["qty_delivered"], it["saldo"], it["unit_cost"], it["in_field_cost"],
                            it["warehouse_status"], it["events_count"]])

    ws = wb.active
    ws.title = "Resumo por Obra"
    ws.append([f"Consolidado de Materiais — todas as obras · Período: {per}" + (f" · Equipe: {team_label}" if team_label else "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Gerado em {_fmt_ts(datetime.now(timezone.utc).isoformat())} · {len(ledgers)} obra(s) com lista de materiais · "
               f"Requisitado/Devolvido/Separado/Entregue = movimentações DENTRO do período · Em campo/Saldo = posição ATUAL"])
    ws.append([])
    cols = ["Código Obra", "Obra", "Equipe", "Localidade", "Status", "Itens", "Qtd. Orçada", "Requisitado (per.)", "Devolvido (per.)", "Separado (per.)", "Entregue (per.)",
            "Em Campo (hoje)", "Saldo a Requisitar (hoje)", "Custo em Campo (R$)", "Custo Orçado (R$)", "Eventos (per.)", "Movimentou no período"]
    hr = header(ws, cols)
    per_obra.sort(key=lambda r: (r[16] != "Sim", r[0]))
    for r in per_obra:
        ws.append(r)
    tot = ["TOTAIS", f"{len(per_obra)} obra(s)", "", "", "", sum(r[5] for r in per_obra), sum(r[6] for r in per_obra), sum(r[7] for r in per_obra), sum(r[8] for r in per_obra),
           sum(r[9] for r in per_obra), sum(r[10] for r in per_obra), sum(r[11] for r in per_obra), sum(r[12] for r in per_obra), sum(r[13] for r in per_obra),
           sum(r[14] for r in per_obra), sum(r[15] for r in per_obra), f"{sum(1 for r in per_obra if r[16] == 'Sim')} sim"]
    ws.append(tot)
    for c in range(1, len(cols) + 1):
        ws.cell(row=ws.max_row, column=c).font = bold
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, min_col=14, max_col=15):
        for cell in row:
            cell.number_format = '#,##0.00'
    widths(ws, [14, 40, 18, 18, 16, 7, 12, 14, 13, 13, 13, 13, 16, 16, 15, 11, 12])
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{ws.max_row - 1}"

    ws2 = wb.create_sheet("Por Material")
    cols2 = ["Código", "Descrição", "Un", "Nº Obras", "Qtd. Orçada (todas)", "Requisitado (per.)", "Devolvido (per.)", "Entregue (per.)", "Em Campo (hoje)", "Saldo a Requisitar (hoje)",
             "Custo Unit. Ref. (R$)", "Custo em Campo (R$)"]
    header(ws2, cols2, 1)
    for m in sorted(por_material.values(), key=lambda x: (-x["field"], x["code"] or x["item"])):
        ws2.append([m["code"], m["item"], m["unit"], len(m["obras"]), round(m["total"], 3), round(m["req"], 3), round(m["ret"], 3), round(m["ent"], 3), round(m["field"], 3),
                    round(m["saldo"], 3), m["unit_cost"], round(m["cost"], 2)])
    ws2.append(["TOTAIS", f"{len(por_material)} material(is)", "", "", round(sum(m['total'] for m in por_material.values()), 3), round(sum(m['req'] for m in por_material.values()), 3),
                round(sum(m['ret'] for m in por_material.values()), 3), round(sum(m['ent'] for m in por_material.values()), 3), round(sum(m['field'] for m in por_material.values()), 3),
                round(sum(m['saldo'] for m in por_material.values()), 3), "", round(sum(m['cost'] for m in por_material.values()), 2)])
    for c in range(1, len(cols2) + 1):
        ws2.cell(row=ws2.max_row, column=c).font = bold
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=11, max_col=12):
        for cell in row:
            cell.number_format = '#,##0.00'
    widths(ws2, [14, 46, 6, 9, 14, 14, 13, 13, 13, 16, 14, 16])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols2))}{max(2, ws2.max_row - 1)}"

    ws3 = wb.create_sheet("Eventos do Período")
    cols3 = ["Data/Hora", "Código Obra", "Obra", "Equipe", "Tipo", "Código", "Material", "Un", "Qtd", "Responsável", "Referência", "Detalhes"]
    header(ws3, cols3, 1)
    for r in eventos:
        ws3.append(r)
    widths(ws3, [17, 14, 34, 16, 26, 14, 40, 6, 9, 22, 18, 70])
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(cols3))}{max(2, ws3.max_row)}"

    ws4 = wb.create_sheet("Posição Atual")
    cols4 = ["Código Obra", "Obra", "Equipe", "Código", "Descrição", "Un", "Origem", "Qtd. Obra", "Requisitado", "Devolvido", "Em Campo", "Separado", "Entregue", "Saldo",
             "Custo Unit. (R$)", "Custo em Campo (R$)", "Status Almox.", "Nº Eventos"]
    header(ws4, cols4, 1)
    for r in posicao:
        ws4.append(r)
    for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row, min_col=15, max_col=16):
        for cell in row:
            cell.number_format = '#,##0.00'
    widths(ws4, [14, 34, 16, 14, 44, 6, 9, 10, 11, 10, 10, 10, 10, 9, 14, 16, 14, 10])
    ws4.freeze_panes = "A2"
    ws4.auto_filter.ref = f"A1:{get_column_letter(len(cols4))}{max(2, ws4.max_row)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
def build_material_ledger_pdf(ledger: dict) -> bytes:
    p = ledger["project"]
    s = ledger["summary"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=10 * mm, bottomMargin=10 * mm, leftMargin=10 * mm, rightMargin=10 * mm,
                            title=f"Movimentacao materiais {p['code']}")
    st = getSampleStyleSheet()
    st.add(ParagraphStyle(name="T", parent=st["Title"], fontSize=16, textColor=colors.HexColor("#050505"), alignment=0, spaceAfter=2))
    st.add(ParagraphStyle(name="S", parent=st["Normal"], fontSize=8.5, textColor=colors.HexColor("#555555"), spaceAfter=6))
    st.add(ParagraphStyle(name="H", parent=st["Heading3"], fontSize=10.5, textColor=colors.HexColor("#050505"), spaceBefore=8, spaceAfter=3))
    st.add(ParagraphStyle(name="C", parent=st["Normal"], fontSize=7.5, leading=9))
    st.add(ParagraphStyle(name="Cb", parent=st["Normal"], fontSize=7.5, leading=9, fontName="Helvetica-Bold"))
    st.add(ParagraphStyle(name="M", parent=st["Normal"], fontSize=7, textColor=colors.HexColor("#888888")))

    def cell(txt, bold=False):
        return Paragraph(str(txt if txt not in (None, "") else "—"), st["Cb"] if bold else st["C"])

    story = [
        Paragraph("HAYAH ENERGY — Relatório de Movimentação de Materiais", st["T"]),
        Paragraph(f"Obra <b>[{p['code']}] {p['name']}</b> · Equipe {p['team_name'] or '—'} · Local {p['location'] or '—'} · "
                  f"Gerado em {_fmt_ts(ledger['generated_at'])}", st["S"]),
    ]
    kpis = [["Itens", "Qtd. Obra", "Requisitado", "Devolvido", "Em Campo", "Entregue", "Saldo a Req.", "Custo Total", "Custo em Campo", "Requisições", "Devoluções", "Eventos"],
            [s["items_count"], _fmt_qty(s["total_qty"]), _fmt_qty(s["requisitioned_qty"]), _fmt_qty(s["returned_qty"]), _fmt_qty(s["in_field_qty"]),
             _fmt_qty(s["delivered_qty"]), _fmt_qty(s["saldo_qty"]), _fmt_brl(s["total_cost"]), _fmt_brl(s["in_field_cost"]), s["requisitions_count"], s["returns_count"], s["events_count"]]]
    kt = Table(kpis, colWidths=[None] * 12)
    kt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#050505")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#E0FF00")),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#050505")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [kt, Spacer(1, 6), Paragraph("Resumo por material", st["H"])]

    rows = [["Código", "Descrição", "Un", "Qtd. Obra", "Requis.", "Devolv.", "Em Campo", "Separado", "Entregue", "Saldo", "Custo Unit.", "Custo Campo", "Eventos"]]
    for it in ledger["items"]:
        rows.append([cell(it["code"]), cell(it["item"] + (" (avulso)" if it["source"] == "avulso" else "")), it["unit"], _fmt_qty(it["qty_total"]), _fmt_qty(it["qty_requisitioned"]),
                     _fmt_qty(it["qty_returned"]), _fmt_qty(it["in_field"]), _fmt_qty(it["qty_separated"]), _fmt_qty(it["qty_delivered"]), _fmt_qty(it["saldo"]),
                     _fmt_brl(it["unit_cost"]), _fmt_brl(it["in_field_cost"]), it["events_count"]])
    rows.append(["", cell("TOTAIS", True), "", _fmt_qty(s["total_qty"]), _fmt_qty(s["requisitioned_qty"]), _fmt_qty(s["returned_qty"]), _fmt_qty(s["in_field_qty"]),
                 _fmt_qty(s["separated_qty"]), _fmt_qty(s["delivered_qty"]), _fmt_qty(s["saldo_qty"]), "", _fmt_brl(s["in_field_cost"]), s["events_count"]])
    t = Table(rows, colWidths=[22 * mm, 74 * mm, 10 * mm, 16 * mm, 16 * mm, 16 * mm, 17 * mm, 17 * mm, 17 * mm, 15 * mm, 20 * mm, 22 * mm, 15 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#050505")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#E0FF00")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (2, 0), (-1, -1), "RIGHT"), ("ALIGN", (2, 0), (2, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F7F7F7")]), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EDEDED")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("TEXTCOLOR", (6, 1), (6, -1), colors.HexColor("#E55400")), ("TEXTCOLOR", (9, 1), (9, -1), colors.HexColor("#005EE0")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#050505")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story += [t, PageBreak(), Paragraph("Histórico de acontecimentos por material", st["H"]),
              Paragraph("Ordem cronológica. Quantidades na unidade do item. Referência = documento da requisição/devolução.", st["M"])]

    ev_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (2, 1), (2, -1), "RIGHT"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ])
    ev_widths = [26 * mm, 42 * mm, 16 * mm, 40 * mm, 30 * mm, 123 * mm]
    for it in ledger["items"]:
        evs = [e for e in ledger["events"] if e["material_id"] == it["id"]]
        story.append(Paragraph(f"{it['code'] + ' · ' if it['code'] else ''}{it['item']} <font color='#888888' size='7'>— {it['unit']} · obra {_fmt_qty(it['qty_total'])} · "
                               f"em campo {_fmt_qty(it['in_field'])} · saldo {_fmt_qty(it['saldo'])}</font>", st["H"]))
        if not evs:
            story.append(Paragraph("Sem movimentações registradas.", st["M"]))
            continue
        rows = [["Data/Hora", "Evento", "Qtd", "Responsável", "Referência", "Detalhes"]]
        for e in evs:
            rows.append([e["ts_label"], cell(e["type_label"]), "" if e["qty"] is None else _fmt_qty(e["qty"]), cell(e["by"]), cell(e["ref"]), cell(e["detail"])])
        et = Table(rows, colWidths=ev_widths, repeatRows=1)
        et.setStyle(ev_style)
        story.append(et)
    general = [e for e in ledger["events"] if not e["material_id"]]
    if general:
        story.append(Paragraph("Eventos gerais da obra (não vinculados a um item)", st["H"]))
        rows = [["Data/Hora", "Evento", "Qtd", "Responsável", "Referência", "Detalhes"]]
        for e in general:
            rows.append([e["ts_label"], cell(e["type_label"]), "" if e["qty"] is None else _fmt_qty(e["qty"]), cell(e["by"]), cell(e["ref"]), cell(e["detail"])])
        gt = Table(rows, colWidths=ev_widths, repeatRows=1)
        gt.setStyle(ev_style)
        story.append(gt)
    story += [Spacer(1, 8), Paragraph("Documento gerado automaticamente pelo sistema Hayah Energy para fins de inventário e controle de materiais.", st["M"])]
    doc.build(story)
    return buf.getvalue()
